"""Tests for the parse_file tool — 解析、权限、路径安全、旧格式与截断。"""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from app.engine.agent.builtin_tools import (
    reset_workspace_context,
    set_workspace_context,
)
from app.engine.agent.parse_tool import _PARSE_TOOLS, PARSE_TOOL_BY_NAME, parse_file
from app.engine.tool.workspace import Workspace
from app.models.file_library import FileConsumerKind, FileRef

# ---------------------------------------------------------------------------
# Fixtures — 临时 workspace + 样本文件工厂
# ---------------------------------------------------------------------------


def _make_workspace(tmp_path: Path, user: str = "user_1", session: str = "sess_1") -> Workspace:
    root = tmp_path / user / session
    ws = Workspace(
        root=root,
        input_dir=root / "input",
        output_dir=root / "output",
        tmp_dir=root / "tmp",
    )
    for d in (ws.input_dir, ws.output_dir, ws.tmp_dir):
        d.mkdir(parents=True, exist_ok=True)
    return ws


@pytest.fixture
def workspace(tmp_path: Path) -> Workspace:
    ws = _make_workspace(tmp_path)
    token = set_workspace_context(ws)
    yield ws
    reset_workspace_context(token)


def _docx_bytes(paragraph: str = "季度营收创新高") -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("财务报告", level=1)
    doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["月份", "营收"])
    ws.append(["1月", 100])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "产品规划"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Contract terms and conditions apply.")
    return doc.tobytes()


# ---------------------------------------------------------------------------
# Tool registration
# ---------------------------------------------------------------------------


def test_tool_registered() -> None:
    assert [t.name for t in _PARSE_TOOLS] == ["parse_file"]
    assert PARSE_TOOL_BY_NAME["parse_file"] is parse_file


# ---------------------------------------------------------------------------
# Happy path — path 入口，四种 Office/PDF 格式
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "factory", "needle", "summary"),
    [
        ("报告.docx", _docx_bytes, "季度营收创新高", "个内容块"),
        ("数据.xlsx", _xlsx_bytes, "月份", "个工作表"),
        ("规划.pptx", _pptx_bytes, "产品规划", "张幻灯片"),
        ("合同.pdf", _pdf_bytes, "Contract terms", "页"),
    ],
)
async def test_parse_by_path(
    workspace: Workspace, filename: str, factory, needle: str, summary: str
) -> None:
    target = workspace.input_dir / filename
    target.write_bytes(factory())

    result = await parse_file.ainvoke({"path": f"input/{filename}"})
    assert not result.startswith("[parse_file]"), result
    assert filename in result
    assert summary in result
    assert needle in result


async def test_parse_output_dir_path(workspace: Workspace) -> None:
    """output/ 下的 agent 自产文件同样可解析。"""
    target = workspace.output_dir / "out.docx"
    target.write_bytes(_docx_bytes())

    result = await parse_file.ainvoke({"path": "output/out.docx"})
    assert "季度营收创新高" in result


async def test_structured_docx_sections(workspace: Workspace) -> None:
    target = workspace.input_dir / "s.docx"
    target.write_bytes(_docx_bytes())

    result = await parse_file.ainvoke({"path": "input/s.docx", "structured": True})
    assert "财务报告" in result


# ---------------------------------------------------------------------------
# file_id 入口 — 归属校验
# ---------------------------------------------------------------------------


def _patch_file_service(monkeypatch, fref: FileRef | None, data: bytes) -> None:
    from app.services.file_service import FileService

    async def fake_load_content(self, file_id: str):
        return None if fref is None else (fref, data)

    monkeypatch.setattr(FileService, "load_content", fake_load_content)


def _fref(owner: str, name: str = "共享.docx") -> FileRef:
    return FileRef(
        id="file_01TEST",
        owner_user_id=owner,
        storage_key=f"{owner}/files/file_01TEST",
        name=name,
        size=1024,
        sha256="0" * 64,
        origin_kind=FileConsumerKind.SESSION_MESSAGE,
        origin_id="sess_1",
    )


async def test_parse_by_file_id(workspace: Workspace, monkeypatch) -> None:
    _patch_file_service(monkeypatch, _fref("user_1"), _docx_bytes())

    result = await parse_file.ainvoke({"file_id": "file_01TEST"})
    assert "季度营收创新高" in result


async def test_file_id_owner_mismatch(workspace: Workspace, monkeypatch) -> None:
    """跨用户 file_id → 拒绝（当前 workspace 属 user_1，文件属 user_2）。"""
    _patch_file_service(monkeypatch, _fref("user_2"), _docx_bytes())

    result = await parse_file.ainvoke({"file_id": "file_01TEST"})
    assert result.startswith("[parse_file]")
    assert "无权访问" in result


async def test_file_id_not_found(workspace: Workspace, monkeypatch) -> None:
    _patch_file_service(monkeypatch, None, b"")

    result = await parse_file.ainvoke({"file_id": "file_missing"})
    assert "不存在" in result


# ---------------------------------------------------------------------------
# 错误分支
# ---------------------------------------------------------------------------


async def test_no_args(workspace: Workspace) -> None:
    result = await parse_file.ainvoke({})
    assert "file_id" in result and "path" in result


async def test_no_workspace(tmp_path: Path) -> None:
    """未设置 workspace 上下文 → 明确报错。"""
    ws = _make_workspace(tmp_path)
    target = ws.input_dir / "x.docx"
    target.write_bytes(_docx_bytes())
    # 不 set_workspace_context（fixture 未启用），确保 ContextVar 为 None。

    result = await parse_file.ainvoke({"path": "input/x.docx"})
    assert "无工作区上下文" in result


async def test_legacy_office_format(workspace: Workspace) -> None:
    target = workspace.input_dir / "old.doc"
    target.write_bytes(b"\xd0\xcf\x11\xe0fake-ole-bytes")

    result = await parse_file.ainvoke({"path": "input/old.doc"})
    assert "旧版二进制格式" in result
    assert ".docx" in result  # 提示另存为新格式


async def test_unsupported_type(workspace: Workspace) -> None:
    target = workspace.input_dir / "app.zip"
    target.write_bytes(b"PK\x03\x04")

    result = await parse_file.ainvoke({"path": "input/app.zip"})
    assert "不支持的文件类型" in result


async def test_path_traversal_blocked(workspace: Workspace) -> None:
    """../ 逃逸出 input/output → 拒绝。"""
    secret = workspace.root.parent / "secret.docx"
    secret.parent.mkdir(parents=True, exist_ok=True)
    secret.write_bytes(_docx_bytes())

    result = await parse_file.ainvoke({"path": "../secret.docx"})
    assert result.startswith("[parse_file]")


async def test_path_outside_input_output(workspace: Workspace) -> None:
    """tmp/ 下文件（在 workspace 内但不在 input/output）→ 拒绝。"""
    target = workspace.tmp_dir / "t.docx"
    target.write_bytes(_docx_bytes())

    result = await parse_file.ainvoke({"path": "tmp/t.docx"})
    assert "越界" in result


async def test_file_missing(workspace: Workspace) -> None:
    result = await parse_file.ainvoke({"path": "input/nope.docx"})
    assert "不存在" in result


# ---------------------------------------------------------------------------
# 截断
# ---------------------------------------------------------------------------


async def test_max_chars_truncation(workspace: Workspace) -> None:
    target = workspace.input_dir / "big.docx"
    target.write_bytes(_docx_bytes(paragraph="很长的内容" * 2000))

    result = await parse_file.ainvoke({"path": "input/big.docx", "max_chars": 1000})
    assert "已截断" in result
    assert len(result) < 3000  # 截断生效（含少量头部/标注开销）


async def test_max_chars_hard_cap(workspace: Workspace) -> None:
    """max_chars 超过 50k 硬上限时被钳制（不实际生成 50k+ 内容，仅验证不报错）。"""
    target = workspace.input_dir / "c.docx"
    target.write_bytes(_docx_bytes())

    result = await parse_file.ainvoke({"path": "input/c.docx", "max_chars": 999999})
    assert "季度营收创新高" in result
