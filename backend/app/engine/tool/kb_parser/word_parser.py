"""Word (.docx) parser — extracts paragraph text via python-docx.

``.doc`` (legacy binary) is NOT supported by python-docx; only ``.docx``.
Word documents have no reliable page concept, so all paragraphs are merged
into a single block with ``page=None``.
"""
from __future__ import annotations

import io

from app.engine.tool.kb_parser.base import ParseResult, TextBlock


def parse_word(file_bytes: bytes) -> ParseResult:
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    # Collect non-empty paragraphs. python-docx does not segment by page.
    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if text:
            parts.append(text)
    # Also extract table cell text (tables often hold key data).
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return ParseResult(blocks=[TextBlock(text="\n".join(parts))], file_type="docx")


__all__ = ["parse_word"]
