"""Excel (.xlsx) parser — extracts cell text from worksheets via openpyxl.

Each worksheet becomes one :class:`TextBlock`. Rows are emitted as pipe-joined
cells (the same convention ``word_parser`` uses for tables) so a sheet reads
as a flat, line-oriented transcript that chunks and retrieves well.

Only values are read (no formula evaluation). ``.xls`` (legacy binary) is NOT
supported by openpyxl; only ``.xlsx``.
"""
from __future__ import annotations

import io

from app.engine.kb.vector.parser.base import ParseResult, TextBlock


def _cell_text(cell) -> str:
    """Best-effort text extraction from a cell (drops None/empty)."""
    v = cell.value
    if v is None:
        return ""
    return str(v).strip()


def parse_xlsx(file_bytes: bytes) -> ParseResult:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    blocks: list[TextBlock] = []
    for idx, ws in enumerate(wb.worksheets, start=1):
        lines: list[str] = []
        for row in ws.iter_rows():
            cells = [_cell_text(c) for c in row]
            # Skip fully-empty rows (keeps the transcript tight).
            if not any(cells):
                continue
            line = " | ".join(c for c in cells if c)
            if line:
                lines.append(line)
        if lines:
            # page = sheet index for citation; section = sheet name.
            blocks.append(
                TextBlock(text="\n".join(lines), page=idx, section=ws.title or "")
            )
    wb.close()
    return ParseResult(
        blocks=blocks, file_type="xlsx", total_pages=len(blocks)
    )


__all__ = ["parse_xlsx"]
